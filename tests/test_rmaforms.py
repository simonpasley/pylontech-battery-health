"""Tests for the RMA form auto-fill.

Logic tests run everywhere. The docx round-trip builds a *synthetic* template
in-memory (we never ship Pylontech's copyrighted files) and is skipped if
python-docx isn't installed.
"""

import io
import os
import tempfile
import unittest

from pylontech import rmaforms as R
from pylontech.diagnose import PackDiagnosis


def _diags():
    # Synthetic test data — fictitious serial numbers, not real hardware.
    return [
        PackDiagnosis(pack_id=1, address=1, barcode="AAA000A00000001", model="US5000",
                      release_date="21-12-27", verdict="HEALTHY", via_master=False,
                      pack_voltage_mv=49260),
        PackDiagnosis(pack_id=4, address=4, barcode="BBB111B11110004", model="US3000C",
                      release_date="21-12-27", verdict="FAILED", via_master=True,
                      pack_voltage_mv=49270,
                      verdict_reasons=["Cell-voltage spread 375 mV exceeds 50 mV failure threshold"]),
        PackDiagnosis(pack_id=9, address=9, barcode="BBB111B11110009", model="US3000C",
                      release_date="21-12-27", verdict="FAILED", via_master=True,
                      pack_voltage_mv=49230,
                      verdict_reasons=["BMS-reported SOH = 0 %"]),
    ]


class TestLogic(unittest.TestCase):
    def test_failed_packs_only(self):
        self.assertEqual([d.address for d in R.failed_packs(_diags())], [4, 9])

    def test_master_and_positions(self):
        d = _diags()
        self.assertEqual(R._master_address(d), 1)
        self.assertEqual(R._position_label(1, 1), "M")
        self.assertEqual(R._position_label(4, 1), "S3")
        self.assertEqual(R._position_label(9, 1), "S8")

    def test_production_date_conversion(self):
        self.assertEqual(R._production_date("21-12-27"), "2021/12/27")
        self.assertEqual(R._production_date(""), "")
        self.assertEqual(R._production_date("garbage"), "garbage")

    def test_fault_description_covers_both_and_event_log(self):
        ev = {4: {"details": ["⚠ 2× under-voltage protection trip — a cell hit its cutoff."]}}
        txt = R.build_fault_description(_diags(), ev)
        self.assertIn("...0004 (pos S3)", txt)
        self.assertIn("...0009 (pos S8)", txt)
        self.assertIn("375 mV", txt)
        self.assertIn("SOH = 0", txt)
        self.assertIn("under-voltage protection trip", txt)   # event-log enrichment

    def test_fault_description_empty_when_no_failures(self):
        healthy = [d for d in _diags() if d.verdict == "HEALTHY"]
        self.assertEqual(R.build_fault_description(healthy), "")


class TestProfileRoundtrip(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self._orig = (R.CONFIG_DIR, R.PROFILE_PATH)
        R.CONFIG_DIR = self.tmp
        R.PROFILE_PATH = os.path.join(self.tmp, "installer.json")

    def tearDown(self):
        R.CONFIG_DIR, R.PROFILE_PATH = self._orig

    def test_save_then_load(self):
        R.save_profile({"name": "Test User", "email": "a@b.c", "bogus": "drop me"})
        prof = R.load_profile()
        self.assertEqual(prof["name"], "Test User")
        self.assertEqual(prof["email"], "a@b.c")
        self.assertNotIn("bogus", prof)          # unknown keys rejected
        self.assertIn("phone", prof)             # all known keys present

    def test_load_missing_file_is_safe(self):
        prof = R.load_profile()
        self.assertEqual(prof["name"], "")


@unittest.skipUnless(R.HAVE_DOCX, "python-docx not installed")
class TestDocxRoundtrip(unittest.TestCase):
    def _template(self):
        """Build a minimal diagnosis-style template in memory."""
        from docx import Document
        doc = Document()
        # contact table
        t = doc.add_table(rows=5, cols=2)
        for i, lbl in enumerate(["Date", "Name", "E-mail", "Phone number", "Shipping address"]):
            t.rows[i].cells[0].text = lbl
        # battery table (header + 3 data rows), 7 cols
        bt = doc.add_table(rows=4, cols=7)
        for i, h in enumerate(["Position", "Serial number", "Model", "Production date",
                               "Battery Alarm (Audible, flashing, steady)",
                               "Blue Residue/Ingress in Ports", "Battery Voltage"]):
            bt.rows[0].cells[i].text = h
        # install table
        it = doc.add_table(rows=4, cols=2)
        for i, lbl in enumerate(["Install Date", "Installation Company",
                                 "Inverter Make and Model", "Installation Address"]):
            it.rows[i].cells[0].text = lbl
        # fault table
        ft = doc.add_table(rows=3, cols=2)
        for i, lbl in enumerate(["Description of Fault", "Fault codes on Inverter",
                                 "Fault codes on BatteryView"]):
            ft.rows[i].cells[0].text = lbl
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def test_fill_diagnosis(self):
        from docx import Document
        profile = {"name": "Test User", "email": "test@example.com",
                   "phone": "07000000000", "address": "1 Test Street",
                   "install_date": "01/01/2023", "company": "Test Installer Ltd",
                   "inverter": "Victron", "fault_codes_inverter": "Internal battery fail"}
        out = R.fill_diagnosis_form(_diags(), profile, {}, template=self._template())
        doc = Document(io.BytesIO(out))

        # contact filled
        contact = doc.tables[0]
        self.assertEqual(contact.rows[1].cells[1].text, "Test User")

        # battery rows: two failed packs, correct positions + serials
        bt = doc.tables[1]
        self.assertEqual(bt.rows[1].cells[0].text, "S3")
        self.assertEqual(bt.rows[1].cells[1].text, "BBB111B11110004")
        self.assertEqual(bt.rows[2].cells[0].text, "S8")
        self.assertIn("confirm on site", bt.rows[1].cells[4].text)
        self.assertEqual(bt.rows[1].cells[6].text, "49.3v")
        # third (leftover) data row cleared
        self.assertEqual(bt.rows[3].cells[0].text, "")

        # fault description present
        ft = doc.tables[3]
        self.assertIn("375 mV", ft.rows[0].cells[1].text)
        self.assertEqual(ft.rows[1].cells[1].text, "Internal battery fail")

    def test_fill_collection(self):
        from docx import Document
        doc0 = Document()
        t = doc0.add_table(rows=6, cols=2)
        for i, lbl in enumerate(["Date", "Battery Model/s", "Serial Number/s",
                                 "Name of Contact on Site", "Contact Number", "Collection Address"]):
            t.rows[i].cells[0].text = lbl
        buf = io.BytesIO(); doc0.save(buf); buf.seek(0)

        out = R.fill_collection_form(_diags(), {"name": "Test User", "phone": "07000000000",
                                                "address": "1 Test Street"}, template=buf)
        doc = Document(io.BytesIO(out))
        t = doc.tables[0]
        self.assertIn("BBB111B11110004", t.rows[2].cells[1].text)
        self.assertIn("BBB111B11110009", t.rows[2].cells[1].text)
        self.assertEqual(t.rows[3].cells[1].text, "Test User")


if __name__ == "__main__":
    unittest.main()
