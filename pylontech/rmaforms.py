"""Auto-fill the Pylontech RMA paperwork from a rack scan.

Two official Pylontech UK forms are needed for a warranty return:

  * Battery Initial Diagnosis Form
  * Battery Collection Request

Both are Word .docx tables. This module fills the fields the tool already
knows (serial numbers, model, production date, pack voltage, and an
auto-written fault description derived from the verdict + event-log summary)
plus a one-time installer profile (name, address, install date, inverter…).

Design notes
------------
* We fill the INSTALLER'S OWN template files, kept in the config dir — we do
  not ship Pylontech's copyrighted, UK-specific templates in the repo. The
  app tells the user where to drop them.
* Fields are located by their row LABEL (left-hand cell), not by position, so
  small template revisions don't desync the fill.
* Merged table cells: python-docx returns a *fresh wrapper* for each column
  of a horizontally-merged cell, so writing col-1 then "clearing" col-2 wipes
  the value you just wrote (they share one underlying <tc>). We therefore only
  ever WRITE the value column and never clear its merge partner.
"""

import json
import os
from datetime import date

from .diagnose import PackDiagnosis

try:
    from docx import Document
    HAVE_DOCX = True
except ImportError:          # python-docx optional at import time
    HAVE_DOCX = False


CONFIG_DIR = os.path.expanduser("~/.pylontech-health")
TEMPLATE_DIR = os.path.join(CONFIG_DIR, "templates")
PROFILE_PATH = os.path.join(CONFIG_DIR, "installer.json")
DIAGNOSIS_TEMPLATE = os.path.join(TEMPLATE_DIR, "diagnosis.docx")
COLLECTION_TEMPLATE = os.path.join(TEMPLATE_DIR, "collection.docx")

# Installer profile fields: (key, label, placeholder). Drives both storage and
# the settings UI so the two never drift apart.
PROFILE_FIELDS = [
    ("name",                  "Name",                     "e.g. Jane Smith"),
    ("company",               "Installation company",     "e.g. ACME Renewables Ltd"),
    ("email",                 "E-mail",                    "e.g. you@example.com"),
    ("phone",                 "Phone number",              "e.g. 07000 000000"),
    ("address",               "Address (shipping / collection / install)", "e.g. 1 Example Street, Town, Postcode"),
    ("install_date",          "Install date",              "e.g. 01/01/2023"),
    ("inverter",              "Inverter make and model",   "e.g. Victron MultiPlus-II"),
    ("fault_codes_inverter",  "Fault code shown on inverter", "e.g. Internal battery fail"),
    ("can_pinout",            "CAN pin configuration note", "e.g. Victron BMS-CAN type A cable"),
    ("signature_name",        "Signature (typed name)",    "Your name"),
]
_PROFILE_DEFAULTS = {k: "" for k, _, _ in PROFILE_FIELDS}


# --------------------------------------------------------------------------- #
# Profile storage
# --------------------------------------------------------------------------- #

def load_profile() -> dict:
    prof = dict(_PROFILE_DEFAULTS)
    try:
        with open(PROFILE_PATH, encoding="utf-8") as f:
            prof.update({k: v for k, v in json.load(f).items() if k in prof})
    except (FileNotFoundError, ValueError):
        pass
    return prof


def save_profile(data: dict) -> dict:
    os.makedirs(CONFIG_DIR, exist_ok=True)
    prof = dict(_PROFILE_DEFAULTS)
    prof.update({k: str(v).strip() for k, v in data.items() if k in prof})
    with open(PROFILE_PATH, "w", encoding="utf-8") as f:
        json.dump(prof, f, indent=2)
    return prof


def status() -> dict:
    """What the UI needs to decide whether form generation is possible."""
    prof = load_profile()
    required = ("name", "email", "phone", "address")
    return {
        "have_docx": HAVE_DOCX,
        "template_dir": TEMPLATE_DIR,
        "diagnosis_template": os.path.isfile(DIAGNOSIS_TEMPLATE),
        "collection_template": os.path.isfile(COLLECTION_TEMPLATE),
        "profile_complete": all(prof.get(k) for k in required),
        "profile": prof,
    }


# --------------------------------------------------------------------------- #
# Low-level docx helpers
# --------------------------------------------------------------------------- #

def _set_cell(cell, text):
    """Write text into a cell, preserving the first run's formatting.

    Never touches a merge partner — see module docstring.
    """
    text = "" if text is None else str(text)
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def _norm(s: str) -> str:
    return " ".join(s.strip().lower().split())


def _fill_label_value_table(table, mapping, value_col=1):
    """For a 2-column 'Label | Value' table, set the value cell for every row
    whose label (normalised) matches a key in `mapping`."""
    norm_map = {_norm(k): v for k, v in mapping.items()}
    for row in table.rows:
        key = _norm(row.cells[0].text)
        for mk, mv in norm_map.items():
            if key.startswith(mk):
                _set_cell(row.cells[value_col], mv)
                break


# --------------------------------------------------------------------------- #
# Data derivation
# --------------------------------------------------------------------------- #

def _master_address(diagnoses) -> int:
    direct = [d.address for d in diagnoses if not d.via_master]
    return min(direct) if direct else min((d.address for d in diagnoses), default=1)


def _position_label(addr: int, master: int) -> str:
    return "M" if addr == master else f"S{addr - master}"


def _production_date(release_date: str) -> str:
    """Convert BMS 'Release Date' (e.g. '21-12-27') to 'YYYY/MM/DD'."""
    parts = (release_date or "").strip().split("-")
    if len(parts) == 3:
        yy, mm, dd = parts
        if len(yy) == 2:
            yy = "20" + yy
        return f"{yy}/{mm}/{dd}"
    return release_date or ""


def failed_packs(diagnoses) -> list:
    return [d for d in diagnoses if d.verdict == "FAILED"]


def build_fault_description(diagnoses, event_summaries=None) -> str:
    """One paragraph describing every FAILED pack, from the verdict reasons
    and (if a log was dumped this session) the event-log summary."""
    event_summaries = event_summaries or {}
    master = _master_address(diagnoses)
    failed = failed_packs(diagnoses)
    if not failed:
        return ""

    bits = []
    for d in failed:
        pos = _position_label(d.address, master)
        tail = d.barcode[-4:] if d.barcode else f"addr{d.address}"
        reasons = "; ".join(d.verdict_reasons) if d.verdict_reasons else "pack reported faulty by BMS"
        clause = f"SN ...{tail} (pos {pos}): {reasons}"
        ev = event_summaries.get(d.address)
        if ev and ev.get("details"):
            evtxt = " ".join(x.lstrip("⚠ ").strip() for x in ev["details"])
            clause += f" Event log: {evtxt}"
        bits.append(clause.rstrip("."))

    header = (f"{len(failed)} pack(s) from this rack failed diagnostics "
              f"(checked with the pylontech console health tool). ")
    return header + ". ".join(bits) + "."


# --------------------------------------------------------------------------- #
# Form fillers
# --------------------------------------------------------------------------- #

def fill_diagnosis_form(diagnoses, profile, event_summaries=None,
                        template=DIAGNOSIS_TEMPLATE):
    """Return filled Diagnosis-form bytes."""
    if not HAVE_DOCX:
        raise RuntimeError("python-docx is not installed")
    doc = Document(template)
    master = _master_address(diagnoses)
    failed = failed_packs(diagnoses)
    today = date.today().strftime("%d/%m/%Y")

    for table in doc.tables:
        labels = {_norm(r.cells[0].text) for r in table.rows}

        if "date" in labels and "name" in labels:                 # contact
            _fill_label_value_table(table, {
                "date": today, "name": profile.get("name", ""),
                "e-mail": profile.get("email", ""),
                "phone number": profile.get("phone", ""),
                "shipping address": profile.get("address", ""),
            })

        elif "install date" in labels:                            # installation
            _fill_label_value_table(table, {
                "install date": profile.get("install_date", ""),
                "installation company": profile.get("company", ""),
                "inverter make and model": profile.get("inverter", ""),
                "installation address": profile.get("address", ""),
            })

        elif "description of fault" in labels:                    # fault info
            _fill_label_value_table(table, {
                "description of fault": build_fault_description(diagnoses, event_summaries),
                "fault codes on inverter": profile.get("fault_codes_inverter", ""),
                "fault codes on batteryview":
                    "n/a - diagnosed via pylontech console health tool; "
                    "event-log + health reports attached",
            })

        elif any(_norm(r.cells[0].text) == "position" for r in table.rows) and len(table.columns) >= 7:
            _fill_battery_table(table, failed, master)

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fill_battery_table(table, failed, master):
    """Fill one row per failed pack in the 7-column battery-info table.

    Physical-inspection columns (alarm / blue residue) can't be known from the
    console, so they get a '(confirm on site)' prompt rather than invented data.
    """
    header_idx = next(i for i, r in enumerate(table.rows)
                      if _norm(r.cells[0].text) == "position")
    data_rows = table.rows[header_idx + 1:]
    for i, row in enumerate(data_rows):
        if i < len(failed):
            d = failed[i]
            vals = [
                _position_label(d.address, master),
                d.barcode,
                d.model,
                _production_date(d.release_date),
                "(confirm on site)",                       # Battery Alarm
                "(confirm on site)",                       # Blue residue / ingress
                f"{d.pack_voltage_mv / 1000:.1f}v",
            ]
            for ci, v in enumerate(vals[:len(row.cells)]):
                _set_cell(row.cells[ci], v)
        else:
            # Clear leftover template placeholder rows (M / S1 / S2 / …).
            for ci in range(len(row.cells)):
                _set_cell(row.cells[ci], "")


def fill_collection_form(diagnoses, profile, template=COLLECTION_TEMPLATE):
    """Return filled Collection-request-form bytes."""
    if not HAVE_DOCX:
        raise RuntimeError("python-docx is not installed")
    doc = Document(template)
    failed = failed_packs(diagnoses)
    models = sorted({d.model for d in failed})
    serials = "   ".join(d.barcode for d in failed)
    today = date.today().strftime("%d/%m/%Y")

    for table in doc.tables:
        _fill_label_value_table(table, {
            "date": today,
            "battery model": (", ".join(models) +
                              (f" (x{len(failed)})" if len(failed) > 1 else "")),
            "serial number": serials,
            "name of contact on site": profile.get("name", ""),
            "contact number": profile.get("phone", ""),
            "collection address": profile.get("address", ""),
        })

    import io
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
